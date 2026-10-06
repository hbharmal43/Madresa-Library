"""Excel export in the same column layout as the original spreadsheet."""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "Barcode",
    "Name of the Book",
    "Category",
    "Author",
    "Publication",
    "Shelf",
    "Status",
    "Borrowed by",
    "Class",
    "Due date",
]


def books_to_xlsx_response(queryset, filename: str | None = None) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Library"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    queryset = queryset.select_related("category", "shelf").order_by("title", "barcode")
    open_loans = {}
    from .models import Loan

    for loan in Loan.objects.open().filter(book__in=queryset).only(
        "book_id", "borrower_name", "borrower_grade", "due_date"
    ):
        open_loans[loan.book_id] = loan

    for book in queryset:
        loan = open_loans.get(book.pk)
        ws.append(
            [
                book.barcode,
                book.title,
                book.category.name if book.category else "",
                book.author,
                book.publication,
                book.shelf.code if book.shelf else "",
                book.get_status_display(),
                loan.borrower_name if loan else "",
                loan.get_borrower_grade_display() if loan and loan.borrower_grade else "",
                loan.due_date.isoformat() if loan else "",
            ]
        )

    widths = [12, 40, 20, 36, 32, 8, 12, 24, 22, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = filename or f"library-{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
