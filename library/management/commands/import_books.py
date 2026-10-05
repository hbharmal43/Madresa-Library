"""
Load books from the library spreadsheet.

    python manage.py import_books data/library_data.xlsx
    python manage.py import_books data/library_data.xlsx --dry-run

Matches on barcode, so re-running updates titles/authors without creating
duplicates and never touches shelf or status of existing books.
"""

from collections import Counter

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from library.models import Book, Category

COLUMN_ALIASES = {
    "barcode": {"barcode", "bar code", "code"},
    "title": {"name of the book", "name", "title", "book"},
    "category": {"category"},
    "author": {"author"},
    "publication": {"publication", "publisher"},
}


def normalise_barcode(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean(value) -> str:
    return "" if value is None else str(value).strip()


class Command(BaseCommand):
    help = "Import books from an .xlsx file with Barcode / Name of the Book / Category / Author / Publication columns."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", default=None, help="Worksheet name (default: first sheet)")
        parser.add_argument("--dry-run", action="store_true", help="Report what would change without saving")

    def handle(self, *args, **options):
        try:
            wb = load_workbook(options["path"], data_only=True, read_only=True)
        except FileNotFoundError as exc:
            raise CommandError(str(exc)) from exc
        ws = wb[options["sheet"]] if options["sheet"] else wb.worksheets[0]

        rows = ws.iter_rows(values_only=True)
        header = [clean(h).lower() for h in next(rows, [])]
        columns = {}
        for key, aliases in COLUMN_ALIASES.items():
            for idx, name in enumerate(header):
                if name in aliases:
                    columns[key] = idx
                    break
        if "barcode" not in columns or "title" not in columns:
            raise CommandError(f"Could not find Barcode and Name of the Book columns. Header was: {header}")

        def cell(row, key):
            idx = columns.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        seen = Counter()
        created = updated = skipped = 0
        problems = []

        with transaction.atomic():
            for line_no, row in enumerate(rows, start=2):
                barcode = normalise_barcode(cell(row, "barcode"))
                title = clean(cell(row, "title"))
                if not barcode and not title:
                    continue
                if not barcode:
                    problems.append(f"row {line_no}: no barcode for '{title}', skipped")
                    skipped += 1
                    continue
                if not title:
                    problems.append(f"row {line_no}: barcode {barcode} has no title, skipped")
                    skipped += 1
                    continue
                seen[barcode] += 1
                if seen[barcode] > 1:
                    problems.append(f"row {line_no}: barcode {barcode} appears again ('{title}'), kept the first one")
                    skipped += 1
                    continue

                category = None
                category_name = clean(cell(row, "category"))
                if category_name:
                    category, _ = Category.objects.get_or_create(name=category_name)

                defaults = {
                    "title": title,
                    "author": clean(cell(row, "author")),
                    "publication": clean(cell(row, "publication")),
                    "category": category,
                }
                book, was_created = Book.objects.update_or_create(barcode=barcode, defaults=defaults)
                created += int(was_created)
                updated += int(not was_created)

            if options["dry_run"]:
                transaction.set_rollback(True)

        for problem in problems:
            self.stdout.write(self.style.WARNING(problem))
        summary = f"{created} created, {updated} updated, {skipped} skipped."
        if options["dry_run"]:
            summary = "DRY RUN (nothing saved): " + summary
        self.stdout.write(self.style.SUCCESS(summary))
